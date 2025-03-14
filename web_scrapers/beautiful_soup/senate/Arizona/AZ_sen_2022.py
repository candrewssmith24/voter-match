import requests
import time
from bs4 import BeautifulSoup
from docx import Document
from requests.exceptions import RequestException

HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/114.0.0.0 Safari/537.36'
}
TIMEOUT = 20
RETRIES = 3
DELAY = 2

def fetch_url(url):
    """Fetch content from a URL with retries."""
    for attempt in range(RETRIES):
        try:
            response = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
            response.raise_for_status()
            return response.text
        except RequestException as e:
            print(f"Attempt {attempt + 1} failed for URL: {url}. Error: {e}")
            if attempt < RETRIES - 1:
                time.sleep(DELAY)
            else:
                print(f"Failed to access {url} after {RETRIES} attempts.")
                return None

def scrape_candidates_and_urls(election_url):
    """Scrape candidate names and URLs from the election page."""
    page_content = fetch_url(election_url)
    if not page_content:
        return {}

    soup = BeautifulSoup(page_content, 'html.parser')
    general_election_heading = soup.find('h5', string="General election for U.S. Senate North Carolina")
    if not general_election_heading:
        print("No 'General election' subsection found.")
        return {}

    results_text = general_election_heading.find_next('p', {'class': 'results_text'})
    if not results_text:
        print("No 'results_text' paragraph found under the general election heading.")
        return {}

    candidate_links = results_text.find_all('a', href=True)
    candidates = {}
    for link in candidate_links:
        candidate_name = link.get_text(strip=True)
        candidate_url = link['href']
        if candidate_name and candidate_url.startswith('https://ballotpedia.org/'):
            candidates[candidate_name] = candidate_url

    return candidates

def scrape_campaign_themes(candidate_url):
    """Scrape campaign themes for a given candidate URL."""
    try:
        response = fetch_url(candidate_url)
        if not response:
            return {}

        soup = BeautifulSoup(response, 'html.parser')

        # Attempt to find a section containing "Campaign themes"
        campaign_themes_section = soup.find(lambda tag: tag.name in ['h2', 'h3'] and 'Campaign themes' in tag.get_text())
        if not campaign_themes_section:
            print(f"No 'Campaign themes' section found for {candidate_url}.")
            return {}

        # Locate the text container after the "Campaign themes" heading
        campaign_content = campaign_themes_section.find_next('div')
        if not campaign_content:
            print(f"No campaign themes content found for {candidate_url}.")
            return {}

        campaign_themes = {}
        current_heading = None

        # Iterate over elements in the campaign content container
        for element in campaign_content.descendants:
            if element.name == 'b':
                current_heading = element.get_text(strip=True)
                campaign_themes[current_heading] = []
            elif element.name == 'p' and current_heading:
                text = element.get_text(strip=True)
                if text:
                    campaign_themes[current_heading].append(text)

        # Join multi-paragraph themes into single strings
        campaign_themes = {key: ' '.join(value) for key, value in campaign_themes.items()}
        return campaign_themes

    except requests.exceptions.RequestException as e:
        print(f"Error accessing the URL {candidate_url}: {e}")
        return {}


def write_to_word(doc, candidate_name, campaign_themes):
    """Write campaign themes to a Word document."""
    doc.add_heading(candidate_name, level=1)
    for theme, text in campaign_themes.items():
        doc.add_heading(theme, level=2)
        doc.add_paragraph(text)

def main():
    election_url = "https://ballotpedia.org/United_States_Senate_election_in_North_Carolina,_2020"
    
    print(f"Scraping candidate URLs from: {election_url}")
    candidates = scrape_candidates_and_urls(election_url)
    if not candidates:
        print("No candidates found.")
        return

    print("\nExtracted Candidate URLs:")
    for name, url in candidates.items():
        print(f"{name}: {url}")

    # Initialize Word document
    doc = Document()

    print("\nScraping campaign themes for each candidate...")
    for candidate_name, candidate_url in candidates.items():
        print(f"\nScraping campaign themes for {candidate_name}...")
        campaign_themes = scrape_campaign_themes(candidate_url)
        if campaign_themes:
            write_to_word(doc, candidate_name, campaign_themes)

    # Save Word document
    output_file = "2020_US_senate_NC_candidate_stances.docx"
    doc.save(output_file)
    print(f"\nCampaign themes have been written to '{output_file}'.")

if __name__ == "__main__":
    main()