import requests
import time
from bs4 import BeautifulSoup
from docx import Document
from requests.exceptions import RequestException

def get_candidate_urls(url):
    response = requests.get(url)
    soup = BeautifulSoup(response.text, 'html.parser')

    candidate_urls = {}
    tables = soup.find_all('table')

    for table in tables:
        for a_tag in table.find_all('a', href=True):
            candidate_name = a_tag.get_text(strip=True)
            candidate_url = a_tag['href']

            if '_presidential_campaign,_2024' in candidate_url.lower() and 'staff' not in candidate_url.lower():
                if not candidate_url.startswith('http'):
                    candidate_url = f"https://ballotpedia.org{candidate_url}"
                candidate_urls[candidate_name] = candidate_url

    print("Filtered Candidate URLs:")
    for name, url in candidate_urls.items():
        print(f"{name}: {url}")

    return candidate_urls

def get_candidate_topics(url):
    response = requests.get(url)
    soup = BeautifulSoup(response.text, 'html.parser')

    policy_section = soup.find('span', {'id': 'Policy_positions'})
    stances = {}
    topics = []

    if policy_section:
        section_content = policy_section.find_parent()
        
        for sibling in section_content.find_next_siblings():
            if sibling.name == 'h2':
                break
            if sibling.name == 'h3':
                topic_title = sibling.get_text(strip=True)
                topics.append(topic_title)
                policy_text = []
                for content in sibling.find_next_siblings():
                    if content.name in ['h2', 'h3']:
                        break
                    if content.name == 'p':
                        policy_text.append(content.get_text(strip=True))
                stances[topic_title] = ' '.join(policy_text)

        print(f"Collected topics for {url}: {topics}")
    else:
        print(f"No 'Policy_positions' section found in {url}")

    return stances, topics

def write_to_word(candidate_name, stances, doc):
    doc.add_heading(candidate_name, level=1)
    for topic, stance in stances.items():
        doc.add_heading(topic, level=2)
        doc.add_paragraph(stance)

def compare_topics(candidate_topics):
    candidates = list(candidate_topics.keys())

    for i in range(len(candidates) - 1):
        first_candidate = candidates[i]
        first_topics = set(candidate_topics[first_candidate].keys())

        for j in range(i + 1, len(candidates)):
            comparison_candidate = candidates[j]
            comparison_topics = set(candidate_topics[comparison_candidate].keys())

            if not comparison_topics:
                continue  # Skip candidates with no policy positions.

            missing_in_comparison = first_topics - comparison_topics
            extra_in_comparison = comparison_topics - first_topics

            if missing_in_comparison or extra_in_comparison:
                print(f"{first_candidate} vs. {comparison_candidate}")
                for topic in missing_in_comparison:
                    print(f"+ {topic}")
                for topic in extra_in_comparison:
                    print(f"- {topic}")
                print()

def find_common_topics(candidate_topics):
    # Extract all candidate topic sets.
    all_topics = [set(stances.keys()) for stances in candidate_topics.values()]
    
    # Find the intersection of all topic sets (common topics).
    common_topics = set.intersection(*all_topics) if all_topics else set()

    if common_topics:
        print("\nCommon Topics Across All Candidates:")
        for topic in sorted(common_topics):
            print(f"- {topic}")
    else:
        print("\nNo common topics found across all candidates.")
        
def main():
    base_url = 'https://ballotpedia.org/Presidential_candidates,_2024'
    candidate_urls = get_candidate_urls(base_url)
    
    candidate_topics = {}
    doc = Document()

    for candidate_name, candidate_url in candidate_urls.items():
        print(f"Processing {candidate_name}...")
        stances, topics = get_candidate_topics(candidate_url)
        if stances:  # Only include candidates who have policy positions.
            candidate_topics[candidate_name] = stances
            write_to_word(candidate_name, stances, doc)

    if len(candidate_topics) < 2:
        print("Not enough candidates with policy positions for comparison.")
        doc.save('C:\\Users\\andia\\Desktop\\GitHub\\voter-match\\scraped_data\\presidential\\2024_presidential_candidate_stances.docx')
        return

    compare_topics(candidate_topics)
    find_common_topics(candidate_topics)
    doc.save('C:\\Users\\andia\\Desktop\\GitHub\\voter-match\\scraped_data\\presidential\\2024_presidential_candidate_stances.docx')
    print("All candidate stances have been written to 2024_presidential_candidate_stances.docx")

if __name__ == '__main__':
    main()